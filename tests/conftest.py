"""conftest — общие фикстуры для тестов docqa."""

import pytest
import docx
from docx.oxml.shared import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def add_hyperlink(paragraph, url: str, text: str) -> None:
    """Добавляет внешнюю гиперссылку в параграф."""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


@pytest.fixture
def clean_doc(tmp_path):
    """Чистый документ без проблем."""
    path = tmp_path / 'clean.docx'
    doc = docx.Document()
    doc.add_heading('Введение', level=1)
    doc.add_paragraph('Нормальный текст достаточной длины.')
    doc.add_heading('Раздел 1', level=2)
    doc.add_paragraph('Ещё один нормальный абзац текста.')
    props = doc.core_properties
    props.title = 'Чистый документ'
    props.author = 'Dmitry Pugachev'
    props.language = 'ru-RU'
    p = doc.add_paragraph('Ссылка: ')
    add_hyperlink(p, 'https://github.com/mityaefimov/docqa', 'DocQA')
    doc.save(str(path))
    return str(path)


@pytest.fixture
def bad_structure_doc(tmp_path):
    """Документ с проблемами структуры."""
    path = tmp_path / 'bad_structure.docx'
    doc = docx.Document()
    doc.add_heading('Глава 1', level=1)
    doc.add_heading('Подглава', level=3)
    doc.add_paragraph('текст между заголовками')
    for _ in range(5):
        doc.add_paragraph('')
    doc.add_paragraph('Х.')
    doc.save(str(path))
    return str(path)


@pytest.fixture
def bad_metadata_doc(tmp_path):
    """Документ без метаданных."""
    path = tmp_path / 'bad_metadata.docx'
    doc = docx.Document()
    doc.add_paragraph('Текст без метаданных достаточной длины.')
    doc.save(str(path))
    return str(path)


@pytest.fixture
def bad_links_doc(tmp_path):
    """Документ с битыми ссылками."""
    path = tmp_path / 'bad_links.docx'
    doc = docx.Document()
    doc.add_paragraph('Обычный текст')
    p = doc.add_paragraph('Ссылки: ')
    add_hyperlink(p, 'github.com/no-scheme', 'без схемы')

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('w:anchor'), 'missing')
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = 'битый якорь'
    run.append(t)
    hyperlink.append(run)
    p._p.append(hyperlink)

    doc.save(str(path))
    return str(path)