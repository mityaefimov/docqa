"""create_test_docs — генератор тестовых документов для docqa."""

import docx
import os
from docx.oxml.shared import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def add_hyperlink(paragraph, url: str, text: str) -> None:
    """Добавляет внешнюю гиперссылку в параграф (через низкоуровневый XML)."""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    hyperlink.append(run)

    paragraph._p.append(hyperlink)


def add_anchor_link(paragraph, anchor: str, text: str) -> None:
    """Добавляет внутреннюю ссылку на закладку (якорь)."""
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('w:anchor'), anchor)

    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    hyperlink.append(run)

    paragraph._p.append(hyperlink)


def set_good_metadata(doc, title: str, author: str) -> None:
    """Устанавливает корректные метаданные."""
    props = doc.core_properties
    props.title = title
    props.author = author
    props.language = 'ru-RU'


def create_document_with_issues(path: str) -> None:
    """Создаёт документ с проблемами структуры, метаданных и ссылок."""
    doc = docx.Document()

    doc.add_heading('Глава 1', level=1)
    doc.add_heading('Подглава', level=3)
    doc.add_paragraph('Это обычный текст между заголовками.')

    for _ in range(5):
        doc.add_paragraph('')

    doc.add_paragraph('Х.')
    doc.add_paragraph('Продолжение главы после короткой строки.')

    doc.add_heading('Глава 2', level=1)
    doc.add_heading('Подглава 2.1', level=2)
    doc.add_paragraph('Обычный текст во второй главе.')

    # битые ссылки: одна без схемы, одна на несуществующую закладку
    p_links = doc.add_paragraph('Ссылки: ')
    add_hyperlink(p_links, 'github.com/mityaefimov/docqa', 'ссылка без схемы')
    add_anchor_link(p_links, 'missing_bookmark', 'битый якорь')

    doc.save(path)
    print(f"✅ Создан: {path}")


def create_clean_document(path: str) -> None:
    """Создаёт чистый документ без проблем."""
    doc = docx.Document()

    doc.add_heading('Введение', level=1)
    doc.add_paragraph('Это введение в документ.')

    doc.add_heading('Раздел 1', level=2)
    doc.add_paragraph('Текст первого раздела с достаточной длиной.')

    doc.add_heading('Подраздел 1.1', level=3)
    doc.add_paragraph('Текст подраздела с нормальной структурой.')

    doc.add_heading('Раздел 2', level=2)
    doc.add_paragraph('Текст второго раздела, тоже нормальный.')

    # хорошая ссылка с корректной схемой
    p_link = doc.add_paragraph('Наш проект: ')
    add_hyperlink(p_link, 'https://github.com/mityaefimov/docqa', 'DocQA на GitHub')

    set_good_metadata(doc, 'Чистый тестовый документ', 'Dmitry Pugachev')

    doc.save(path)
    print(f"✅ Создан: {path}")


def create_document_with_orphans(path: str) -> None:
    """Создаёт документ с короткими одинокими строками."""
    doc = docx.Document()

    doc.add_heading('Глава с артефактами', level=1)
    doc.add_paragraph('Нормальный текст в начале.')

    doc.add_paragraph('1.')
    doc.add_paragraph('а')
    doc.add_paragraph('-')

    doc.add_paragraph('После артефактов идёт нормальный текст достаточной длины.')

    doc.core_properties.author = 'Dmitry Pugachev'

    doc.save(path)
    print(f"✅ Создан: {path}")


def main():
    """Создаёт все тестовые документы в папке проекта."""
    project_dir = os.path.dirname(os.path.abspath(__file__))

    print("🔨 Создаю тестовые документы...")
    print()

    create_document_with_issues(os.path.join(project_dir, 'test_sample.docx'))
    create_clean_document(os.path.join(project_dir, 'test_clean.docx'))
    create_document_with_orphans(os.path.join(project_dir, 'test_orphans.docx'))

    print()
    print("🎉 Готово! Три тестовых документа созданы.")


if __name__ == "__main__":
    main()