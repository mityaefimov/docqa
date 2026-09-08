"""structure — проверка структурной целостности документа."""

import docx
from typing import List

from .base import BaseValidator, Issue


class StructureValidator(BaseValidator):
    """Проверяет структуру документа:
    
    - иерархию заголовков (нет прыжков с H1 на H3)
    - избыточные пустые параграфы подряд
    - подозрительно короткие одинокие строки
    """

    def validate(self) -> List[Issue]:
        """Запускает все проверки структуры."""
        doc = docx.Document(self.doc_path)
        self._check_headings(doc)
        self._check_empty_paragraphs(doc)
        self._check_orphaned_text(doc)
        return self.issues

    def _check_headings(self, doc) -> None:
        """Проверяет, что уровни заголовков не пропускаются."""
        prev_level = 0
        for i, para in enumerate(doc.paragraphs):
            if 'Heading' not in para.style.name:
                continue
            try:
                level = int(para.style.name.split()[-1])
            except (ValueError, IndexError):
                continue

            if prev_level > 0 and level > prev_level + 1:
                self._add_issue(
                    severity='warning',
                    message=f'Пропущен уровень: после H{prev_level} сразу H{level}',
                    location=f'Paragraph {i + 1}'
                )
            prev_level = level

    def _check_empty_paragraphs(self, doc) -> None:
        """Находит избыточные пустые параграфы подряд (больше 3)."""
        empty_count = 0
        start_idx = 0

        for i, para in enumerate(doc.paragraphs):
            if not para.text.strip():
                if empty_count == 0:
                    start_idx = i
                empty_count += 1
            else:
                if empty_count > 3:
                    self._add_issue(
                        severity='info',
                        message=f'{empty_count} пустых параграфов подряд',
                        location=f'Paragraphs {start_idx + 1}-{i}'
                    )
                empty_count = 0

    def _check_orphaned_text(self, doc) -> None:
        """Находит подозрительно короткие одинокие строки (меньше 3 символов)."""
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if 0 < len(text) < 3 and 'Heading' not in para.style.name:
                self._add_issue(
                    severity='warning',
                    message=f'Подозрительно короткий текст: "{text}"',
                    location=f'Paragraph {i + 1}'
                )